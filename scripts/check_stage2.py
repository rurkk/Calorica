#!/usr/bin/env python3
"""Structural verification, complementary to mandatory visual page inspection."""
import argparse
import hashlib
import json
import re
from pathlib import Path
from zipfile import ZipFile

from docx import Document
from lxml import etree
from pypdf import PdfReader

from build_stage2 import Builder, ROOT

parser = argparse.ArgumentParser()
parser.add_argument("pdf", type=Path)
args = parser.parse_args()
docx_path = ROOT / "docs/stage-2/Calorica-stage-2-technical-specification.docx"
doc = Document(docx_path)
reader = PdfReader(args.pdf)
pdf_text = "\n".join(page.extract_text() for page in reader.pages)
doc_text = "\n".join(p.text for p in doc.paragraphs)
doc_text += "\n" + "\n".join(c.text for t in doc.tables for row in t.rows for c in row.cells)
errors = []
counts = {}
for name, glob, pattern in [
    ("FR", "docs/requirements/functional/*.md", r"\*\*(FR-[A-Z]+-\d+[a-z]?)\."),
    ("NFR", "docs/requirements/nonfunctional/*.md", r"^\| (NFR-[A-Z]+-\d+) \|"),
    ("AC", "docs/requirements/acceptance/*.md", r"^\| (AC-\d+) \|"),
]:
    ids = [m for file in ROOT.glob(glob) for m in re.findall(pattern, file.read_text(), re.M)]
    counts[name] = len(ids)
    if len(ids) != len(set(ids)): errors.append(f"Duplicate source IDs: {name}")
    for ident in ids:
        for label, text in [("DOCX", doc_text), ("PDF", pdf_text)]:
            if not re.search(r"\b" + re.escape(ident) + r"\b", text):
                errors.append(f"Missing {ident} in {label}")

meta = json.loads((ROOT / "build/stage-2/assembly.json").read_text())
for path, digest in meta["sources"].items():
    if hashlib.sha256((ROOT / path).read_bytes()).hexdigest() != digest:
        errors.append(f"Changed source since build: {path}")

builder = Builder(args.pdf)
if builder.toc_pages != meta["toc_pages"]:
    errors.append("TOC page numbers differ from latest render")
if len(builder.toc_pages) != len(builder.section_titles):
    errors.append("Missing major heading in rendered PDF")
toc_text = reader.pages[1].extract_text()
for title, page in builder.toc_pages.items():
    if not re.search(re.escape(title) + r"\.+\s*" + str(page) + r"\b", toc_text):
        errors.append(f"Wrong cached TOC entry: {title}")

with ZipFile(docx_path) as z:
    styles = etree.fromstring(z.read("word/styles.xml"))
    document = etree.fromstring(z.read("word/document.xml"))
    ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
    if styles.xpath(".//w:pBdr", namespaces=ns): errors.append("Unexpected paragraph border")
    if not document.xpath(".//w:instrText[contains(text(), 'TOC')]", namespaces=ns):
        errors.append("No native TOC field")
    if document.xpath(".//w:ins | .//w:del | .//w:commentRangeStart", namespaces=ns):
        errors.append("Unexpected revision markup")
    if not document.xpath(".//w:tblHeader", namespaces=ns): errors.append("No repeated table headers")

for token in ["[[TOC]]", "TODO", "TBD", "turn26view", "Lorem ipsum"]:
    if token in doc_text or token in pdf_text: errors.append(f"Unresolved marker: {token}")

report = {
    "pages": len(reader.pages), "requirements": counts,
    "toc_pages": builder.toc_pages, "source_count": len(meta["sources"]),
    "docx_sha256": hashlib.sha256(docx_path.read_bytes()).hexdigest(),
    "pdf_sha256": hashlib.sha256(args.pdf.read_bytes()).hexdigest(),
    "errors": errors,
    "visual_review": "Separate visual inspection is required; this script does not perform it.",
}
print(json.dumps(report, ensure_ascii=False, indent=2))
raise SystemExit(bool(errors))
