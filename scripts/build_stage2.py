#!/usr/bin/env python3
"""Build the stage 2 DOCX from versioned Markdown; run with bundled Python.

Only document artifacts and build metadata are generated. Source MD is read-only.
After rendering, --toc-pdf seeds the TOC field with verified page numbers.
"""
import argparse
import hashlib
import json
import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor
from pypdf import PdfReader

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs/stage-2"
REQ = "docs/requirements/"
ARCH = "docs/architecture/"
SECTIONS = [
    ("Общие сведения", [
        ("Основание и сокращения", "docs/stage-2/document-introduction.md", None),
        ("Назначение продукта и команда", REQ + "vision/product-vision.md", None),
        ("Пользователи и проблема", REQ + "vision/target-users-and-problem.md", None),
        ("Термины предметной области", REQ + "README.md", "## Термины"),
    ]),
    ("Объём и план на семестр", [(None, REQ + "planning/semester-plan.md", None)]),
    ("Функциональные требования", [
        ("Аутентификация и роли", REQ + "functional/authentication-and-roles.md", None),
        ("Каталог продуктов", REQ + "functional/product-catalog.md", None),
        ("Дневник питания", REQ + "functional/meal-diary.md", None),
        ("Аналитика", REQ + "functional/analytics.md", None),
    ]),
    ("Правила предметной области", [
        ("Граф состава продуктов", REQ + "domain/product-composition-graph.md", None),
        ("Выход готового блюда", REQ + "domain/cooking-yield.md", None),
        ("Единицы измерения", REQ + "domain/units-of-measurement.md", None),
        ("Лейблы продуктов и приёмов пищи", REQ + "domain/product-and-meal-labels.md", None),
        ("Права доступа и мягкое удаление", REQ + "domain/access-and-soft-deletion.md", None),
        ("Принципы расчётов и хранения", REQ + "domain/calculation-principles.md", None),
        ("Стартовый базовый каталог", REQ + "data/base-catalog.md", None),
    ]),
    ("Нефункциональные требования", [
        ("Общие положения", REQ + "nonfunctional/README.md", "## Границы"),
        ("Производительность и нагрузка", REQ + "nonfunctional/performance.md", None),
        ("Надёжность и эксплуатация", REQ + "nonfunctional/reliability.md", None),
        ("Безопасность и изоляция данных", REQ + "nonfunctional/security.md", None),
        ("Совместимость и удобство", REQ + "nonfunctional/usability-and-compatibility.md", None),
    ]),
    ("Бизнес процессы", [(None, REQ + "processes/as-is-to-be.md", "## Участники, начало и результат")]),
    ("Технические ограничения", [
        ("Архитектура", ARCH + "README.md", "## Компоненты и границы"),
        ("Сессии время редактирование и кэш", ARCH + "runtime-policies.md", None),
        ("Точность чисел", ARCH + "numeric-precision.md", None),
    ]),
    ("Порядок приёмки приложения", [(None, REQ + "acceptance/README.md", "## Набор данных для функциональной проверки")]),
    ("Комплект поставки и порядок изменений", [(None, "docs/stage-2/delivery.md", None)]),
]

def element(tag, **attrs):
    el = OxmlElement(tag)
    for key, value in attrs.items():
        el.set(qn("w:" + key), str(value))
    return el

def heading_text(s):
    s = re.sub(r"^\d+\.\s*", "", s)
    return re.sub(r"\s+", " ", re.sub(r"[^\w\s]", " ", s)).strip()

def field(p, instruction, display):
    p.add_run()._r.append(element("w:fldChar", fldCharType="begin"))
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " " + instruction + " "
    p.add_run()._r.append(instr)
    p.add_run()._r.append(element("w:fldChar", fldCharType="separate"))
    p.add_run(display)
    p.add_run()._r.append(element("w:fldChar", fldCharType="end"))

class Builder:
    def __init__(self, toc_pdf=None):
        self.doc = Document()
        self.sources = {}
        self.bibliography = {}
        self.links = {}
        self.section_titles = [x[0] for x in SECTIONS] + ["Источники"]
        self.toc_pages = self.read_pages(toc_pdf) if toc_pdf else {}
        for i, (_, items) in enumerate(SECTIONS, 1):
            for j, (title, path, _) in enumerate(items, 1):
                self.links[str((ROOT / path).resolve())] = f"{i}.{j}" if title else str(i)
        self.links[str((ROOT / (REQ + "decisions/README.md")).resolve())] = "9"
        self.setup()

    def read_pages(self, pdf):
        result = {}
        for i, page in enumerate(PdfReader(pdf).pages, 1):
            if i <= 2:
                continue
            lines = [re.sub(r"\s+", " ", line).strip() for line in page.extract_text().splitlines()]
            for title in self.section_titles:
                if any(re.fullmatch(r"\d+\s+" + re.escape(title), x) for x in lines):
                    result.setdefault(title, i)
        return result

    def setup(self):
        doc = self.doc
        sec = doc.sections[0]
        sec.page_width, sec.page_height = Mm(210), Mm(297)
        sec.left_margin, sec.right_margin = Mm(30), Mm(15)
        sec.top_margin = sec.bottom_margin = Mm(20)
        sec.footer_distance = Mm(10)
        sec.different_first_page_header_footer = True
        normal = doc.styles["Normal"]
        normal.font.name, normal.font.size = "Times New Roman", Pt(12)
        normal.font.color.rgb = RGBColor(0, 0, 0)
        normal.paragraph_format.line_spacing = 1.15
        normal.paragraph_format.space_after = Pt(5)
        normal.paragraph_format.widow_control = True
        normal.element.get_or_add_rPr().append(element("w:lang", val="ru-RU"))
        for name, size in [("Title", 22), ("Subtitle", 16), ("Heading 1", 16), ("Heading 2", 14), ("Heading 3", 12), ("Heading 4", 12)]:
            style = doc.styles[name]
            style.font.name, style.font.size = "Times New Roman", Pt(size)
            style.font.color.rgb = RGBColor(0, 0, 0)
            style.font.bold = name != "Subtitle"
            style.paragraph_format.line_spacing = 1.15
            style.paragraph_format.space_before = Pt(14)
            style.paragraph_format.space_after = Pt(7)
            style.paragraph_format.keep_with_next = True
        for style in doc.styles:
            for border in style.element.xpath(".//w:pBdr"):
                border.getparent().remove(border)
            for spacing in style.element.xpath(".//w:contextualSpacing"):
                spacing.getparent().remove(spacing)
            for fonts in style.element.xpath(".//w:rFonts"):
                for attr in list(fonts.attrib):
                    if "Theme" in attr: del fonts.attrib[attr]
                for attr in ("ascii", "hAnsi", "eastAsia", "cs"):
                    fonts.set(qn("w:" + attr), "Times New Roman")
        numbering = doc.part.numbering_part.element
        abstract = element("w:abstractNum", abstractNumId=80)
        abstract.append(element("w:multiLevelType", val="multilevel"))
        for level in range(2):
            lvl = element("w:lvl", ilvl=level)
            for el in [element("w:start", val=1), element("w:numFmt", val="decimal"),
                       element("w:pStyle", val=f"Heading{level+1}"),
                       element("w:lvlText", val="%1" if level == 0 else "%1.%2"),
                       element("w:suff", val="space")]:
                lvl.append(el)
            abstract.append(lvl)
        numbering.append(abstract)
        num = element("w:num", numId=80)
        num.append(element("w:abstractNumId", val=80))
        numbering.append(num)
        for level in range(2):
            numpr = element("w:numPr")
            numpr.append(element("w:ilvl", val=level))
            numpr.append(element("w:numId", val=80))
            doc.styles[f"Heading {level+1}"].element.get_or_add_pPr().append(numpr)
        footer = sec.footer.paragraphs[0]
        footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
        field(footer, "PAGE", "1")
        for run in footer.runs:
            run.font.size = Pt(11)
        doc.core_properties.title = "Техническое задание Calorica"
        doc.core_properties.author = "Команда Calorica"
        doc.core_properties.subject = "Второй этап учебного проекта"
        doc.core_properties.keywords = "Calorica, техническое задание, этап 2"

    def cover(self):
        p = self.doc.add_paragraph("Университет ИТМО")
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(65)
        p = self.doc.add_paragraph("Техническое задание", "Title")
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p = self.doc.add_paragraph("Calorica", "Title")
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p = self.doc.add_paragraph("Мобильное приложение для учёта питания,\nкалорий и БЖУ с серверной частью")
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(24)
        p = self.doc.add_paragraph("Второй этап учебного проекта\nВерсия 0.2 · 28 сентября 2026 года")
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(30)
        self.doc.add_paragraph("Команда проекта").runs[0].bold = True
        self.doc.add_paragraph("Валера — DevOps-инженер\nКамила — тестировщик\nБогдан и Николай — разработчики\nРауль и Мария — аналитики")
        p = self.doc.add_paragraph("2026")
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        self.doc.add_page_break()

    def toc(self):
        self.doc.add_paragraph("Содержание", "Title")
        # A native TOC field with populated cached entries. Word can update it.
        for i, title in enumerate(self.section_titles, 1):
            p = self.doc.add_paragraph()
            p.paragraph_format.line_spacing = 1.15
            p.paragraph_format.space_after = Pt(12)
            p.paragraph_format.tab_stops.add_tab_stop(Mm(164), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
            if i == 1:
                p.add_run()._r.append(element("w:fldChar", fldCharType="begin"))
                instr = OxmlElement("w:instrText")
                instr.set(qn("xml:space"), "preserve")
                instr.text = ' TOC \\o "1-1" \\h \\z '
                p.add_run()._r.append(instr)
                p.add_run()._r.append(element("w:fldChar", fldCharType="separate"))
            p.add_run(f"{i} {title}\t{self.toc_pages.get(title, '')}")
            if i == len(self.section_titles):
                p.add_run()._r.append(element("w:fldChar", fldCharType="end"))
        self.doc.add_page_break()

    def inline(self, p, text, source):
        def link(m):
            label, target = m.group(1), m.group(2)
            if target.startswith("http"):
                if target not in self.bibliography:
                    self.bibliography[target] = (len(self.bibliography) + 1, label)
                return f"{label} [{self.bibliography[target][0]}]"
            ref = self.links.get(str((source.parent / target.split("#")[0]).resolve()))
            if not ref:
                raise ValueError(f"Unmapped link {source}: {target}")
            return f"{label} (раздел {ref})"
        text = re.sub(r"\[([^\]]+)\]\(([^)]+)\)", link, text)
        for frag in re.split(r"(\*\*.*?\*\*|`[^`]+`)", text):
            if frag.startswith("**"):
                p.add_run(frag[2:-2]).bold = True
            elif frag.startswith("`"):
                run = p.add_run(frag[1:-1])
                run.font.name = "Consolas"
                run.font.size = Pt(11)
            else:
                p.add_run(frag)

    def table(self, rows, source):
        # Long requirement and acceptance prose reads better as numbered records.
        if rows[0][0] in ("ID", "ID проверки"):
            for row in rows[1:]:
                p = self.doc.add_paragraph()
                p.paragraph_format.keep_with_next = True
                p.add_run(row[0]).bold = True
                if row[0].startswith("AC-"):
                    p.add_run(" · " + row[1])
                    body = row[2]
                    self.inline(self.doc.add_paragraph(), body, source)
                else:
                    self.inline(p, ". " + row[1], source)
                    p = self.doc.add_paragraph()
                    p.add_run("Проверка. ").italic = True
                    self.inline(p, row[2], source)
            return
        t = self.doc.add_table(rows=0, cols=len(rows[0]))
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        t.autofit = False
        if rows[0][0] == "Недели": widths = [18, 75, 72]
        elif len(rows[0]) == 3: widths = [43, 68, 54]
        elif rows[0][0] == "Группа": widths = [57, 108]
        else: widths = [64, 101]
        for col, width in zip(t.columns, widths): col.width = Mm(width)
        borders = element("w:tblBorders")
        for side in ["top", "left", "bottom", "right", "insideH", "insideV"]:
            borders.append(element("w:" + side, val="single", sz=4, color="D9D9D9"))
        t._tbl.tblPr.append(borders)
        for i, values in enumerate(rows):
            row = t.add_row()
            row._tr.get_or_add_trPr().append(element("w:cantSplit"))
            if i == 0: row._tr.get_or_add_trPr().append(element("w:tblHeader"))
            for cell, value, width in zip(row.cells, values, widths):
                cell.width = Mm(width)
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                pr = cell._tc.get_or_add_tcPr()
                margins = element("w:tcMar")
                for side in ["top", "left", "bottom", "right"]:
                    margins.append(element("w:" + side, w=90, type="dxa"))
                pr.append(margins)
                if i == 0: pr.append(element("w:shd", fill="E8EDF2"))
                p = cell.paragraphs[0]
                p.paragraph_format.line_spacing = 1.1
                p.paragraph_format.space_after = Pt(3)
                p.paragraph_format.space_before = Pt(3)
                self.inline(p, value, source)
                for run in p.runs:
                    run.font.size = Pt(11)
                    if i == 0: run.bold = True
        self.doc.add_paragraph().paragraph_format.space_after = Pt(0)

    def markdown(self, text, source, base=3):
        lines = text.strip().splitlines()
        i = 0
        while i < len(lines):
            line = lines[i].strip()
            if not line:
                i += 1; continue
            if line.startswith("```"):
                i += 1
                while i < len(lines) and not lines[i].startswith("```"):
                    p = self.doc.add_paragraph()
                    p.paragraph_format.line_spacing = 1.1
                    p.paragraph_format.space_after = Pt(3)
                    run = p.add_run(lines[i])
                    run.font.name, run.font.size = "Consolas", Pt(10)
                    i += 1
                i += 1; continue
            if line.startswith("|"):
                rows = []
                while i < len(lines) and lines[i].strip().startswith("|"):
                    vals = [x.strip() for x in lines[i].strip().strip("|").split("|")]
                    if not all(re.fullmatch(r"[:\- ]+", x) for x in vals): rows.append(vals)
                    i += 1
                self.table(rows, source); continue
            match = re.match(r"^(#+)\s+(.*)", line)
            if match:
                level = min(4, base + len(match[1]) - 2)
                self.doc.add_heading(heading_text(match[2]), level=level)
                i += 1; continue
            item = re.match(r"^(- |\d+\. )(.*)", line)
            content = item[2] if item else line
            prefix = item[1] if item else ""
            i += 1
            while i < len(lines) and lines[i].strip() and not re.match(r"^\s*(#|\||```|- |\d+\. )", lines[i]):
                content += " " + lines[i].strip(); i += 1
            p = self.doc.add_paragraph()
            if item:
                p.paragraph_format.left_indent = Mm(5)
                p.paragraph_format.first_line_indent = Mm(-5)
                p.add_run("• " if prefix == "- " else prefix)
            self.inline(p, content, source)

    def load(self, path, start):
        source = ROOT / path
        raw = source.read_text()
        self.sources[path] = hashlib.sha256(raw.encode()).hexdigest()
        text = raw.split("\n", 1)[1]
        if start:
            text = raw[raw.index(start):]
            if path.endswith("requirements/README.md") or path.endswith("nonfunctional/README.md"):
                text = text.split("\n", 1)[1]
        if path.endswith("semester-plan.md"):
            text = "План рассчитан на 12 относительных недель от старта реализации. Календарные даты уточняются по расписанию курса и доступности команды.\n\n" + raw[raw.index("## Обязательный результат"):]
        if path.endswith("reliability.md"):
            text = text.replace("Документ не подтверждает наличие мониторинга или резервных копий сейчас.", "")
            text = text.replace("документ не разрешает покупать сервис или менять существующий сервер автоматически.", "изменения инфраструктуры согласуются отдельно.")
        if path.endswith("numeric-precision.md"):
            text = text.replace("Ниже выбран технический ориентир по поручению владельца; это не утверждение о\nвнутреннем хранении данных у конкурентов.", "Публичные примеры не устанавливают внутреннюю схему хранения данных конкурентов.")
        return source, text

    def build(self):
        self.cover(); self.toc()
        for index, (title, items) in enumerate(SECTIONS):
            self.doc.add_heading(title, level=1)
            if index == 4:
                self.doc.add_paragraph("Показатели ниже задают условия приёмки учебного MVP. Испытания выполняются после реализации; наличие требования не подтверждает его выполнение.")
            if index == 7:
                self.doc.add_paragraph("Приёмка приложения проводится после реализации. Для каждого сценария фиксируются проверенная версия, результат и доказательство. Проверка документации второго этапа не заменяет эти испытания.")
            for subtitle, path, start in items:
                if subtitle: self.doc.add_heading(subtitle, level=2)
                source, content = self.load(path, start)
                self.markdown(content, source, base=3)
        p = self.doc.add_heading("Источники", level=1)
        p.paragraph_format.page_break_before = True
        self.doc.add_paragraph("Технические источники используются для обоснования решений. Требования Calorica и выбранные целевые значения определены в настоящем ТЗ.")
        for url, (num, title) in self.bibliography.items():
            p = self.doc.add_paragraph(f"[{num}] {title}.\n{url}")
            p.paragraph_format.line_spacing = 1.1
            for run in p.runs: run.font.size = Pt(11)
        OUT.mkdir(parents=True, exist_ok=True)
        output = OUT / "Calorica-stage-2-technical-specification.docx"
        self.doc.save(output)
        meta = {"sources": self.sources, "toc_pages": self.toc_pages, "sources_count": len(self.sources), "bibliography": self.bibliography}
        (ROOT / "build/stage-2").mkdir(parents=True, exist_ok=True)
        (ROOT / "build/stage-2/assembly.json").write_text(json.dumps(meta, ensure_ascii=False, indent=2))
        print(output)

if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--toc-pdf", type=Path)
    args = parser.parse_args()
    Builder(args.toc_pdf).build()
