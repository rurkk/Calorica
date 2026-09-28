"""Build only the concise semester MVP specification; leave earlier documents intact."""
import argparse
import re
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[3]
HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(ROOT / "scripts"))
from build_stage2 import Builder
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Mm, Pt
from pypdf import PdfReader

SOURCE = HERE / "technical-specification.md"
OUTPUT = HERE / "Calorica-semester-MVP-TZ.docx"


class SemesterBuilder(Builder):
    def table(self, rows, source):
        super().table(rows, source)
        if rows[0] == ["Работа", "Человеко-дни"]:
            table = self.doc.tables[-1]
            for column, width in zip(table.columns, (137, 28)):
                column.width = Mm(width)
            for row in table.rows:
                for cell, width in zip(row.cells, (137, 28)):
                    cell.width = Mm(width)
                for paragraph in row.cells[1].paragraphs:
                    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER


def blocks():
    parts = re.split(r"^## (.+)$", SOURCE.read_text(), flags=re.M)
    return list(zip(parts[1::2], parts[2::2]))


def build():
    b = SemesterBuilder()
    d = b.doc
    d.core_properties.title = "Техническое задание Calorica Семестровый MVP"
    d.core_properties.subject = "Второй этап Семестровый MVP"
    p = d.add_paragraph("Университет ИТМО")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(68)
    for text in ("Техническое задание", "Calorica"):
        p = d.add_paragraph(text, "Title")
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p = d.add_paragraph("Мобильное приложение для учёта питания\nс серверной частью")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(25)
    p = d.add_paragraph("Семестровый MVP\nВторой этап учебного проекта\nВерсия 1.0 · 28 сентября 2026 года")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(28)
    d.add_paragraph("Команда проекта и формальные роли").runs[0].bold = True
    d.add_paragraph("Валера — DevOps-инженер\nКамила — тестировщик\nБогдан и Николай — разработчики\nРауль и Мария — аналитики")
    p = d.add_paragraph("2026")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for title, content in blocks():
        heading = d.add_heading(title, level=1)
        heading.paragraph_format.page_break_before = True
        b.markdown(content, SOURCE, base=2)
    d.save(OUTPUT)
    print(OUTPUT)


def check():
    source = SOURCE.read_text()
    doc = Document(OUTPUT)
    doc_text = "\n".join(p.text for p in doc.paragraphs)
    doc_text += "\n" + "\n".join(c.text for t in doc.tables for r in t.rows for c in r.cells)
    pdf = PdfReader(OUTPUT.with_suffix(".pdf"))
    pdf_text = "\n".join(p.extract_text() for p in pdf.pages)
    expected = {f"SM-{i:02}" for i in range(1, 10)} | {f"SB-{i:02}" for i in range(1, 7)} | {f"SN-{i:02}" for i in range(1, 9)}
    errors = []
    for label, text in (("MD", source), ("DOCX", doc_text), ("PDF", pdf_text)):
        ids = re.findall(r"(?:SM|SB|SN)-\d{2}", text)
        if not expected.issubset(ids): errors.append(f"{label}: missing IDs {expected - set(ids)}")
    def norm(text):
        return re.sub(r"\s+", "", text.replace("**", "").replace("\u00ad", ""))
    for title, content in blocks():
        if norm(title) not in norm(doc_text) or norm(title) not in norm(pdf_text):
            errors.append(f"Missing heading: {title}")
        for paragraph in re.split(r"\n\s*\n", content.strip()):
            if paragraph.startswith("###"):
                heading = paragraph.lstrip("# ")
                if norm(heading) not in norm(doc_text) or norm(heading) not in norm(pdf_text):
                    errors.append(f"Missing subheading: {heading}")
                continue
            if paragraph.startswith("|"):
                for row in paragraph.splitlines():
                    for cell in row.strip("|").split("|"):
                        cell = cell.strip()
                        if not cell or re.fullmatch(r"[-: ]+", cell):
                            continue
                        if norm(cell) not in norm(doc_text) or norm(cell) not in norm(pdf_text):
                            errors.append(f"Missing table cell: {cell}")
                continue
            for text in (doc_text, pdf_text):
                # Lists are several independent paragraphs in the output.
                for item in re.split(r"\n- ", paragraph):
                    item = re.sub(r"^- ", "", item)
                    if norm(item) not in norm(text):
                        errors.append(f"Text mismatch: {item[:65]}")
    if doc.element.xpath(".//w:ins | .//w:del | .//w:pBdr"):
        errors.append("Unexpected revisions or borders")
    if len(pdf.pages) != 7:
        errors.append(f"Expected 7 pages, got {len(pdf.pages)}")
    print({"pages": len(pdf.pages), "SM": 9, "SB": 6, "SN": 8, "errors": errors})
    if errors: raise SystemExit(1)


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--check", action="store_true")
    args = parser.parse_args()
    check() if args.check else build()
